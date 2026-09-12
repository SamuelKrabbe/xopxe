"""
Gera o arquitetura.docx a partir do arquitetura.md, trocando o bloco
```mermaid pelo diagrama-de-classes.png.

    pip install python-docx
    cd docs && python3 gerar-docx.py

Os três caminhos podem ser passados como argumentos, nesta ordem:
markdown, imagem e arquivo de saída.
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))

SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(AQUI, "arquitetura.md")
IMG = sys.argv[2] if len(sys.argv) > 2 else os.path.join(AQUI, "diagrama-de-classes.png")
DST = sys.argv[3] if len(sys.argv) > 3 else os.path.join(AQUI, "arquitetura.docx")

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

code_style = doc.styles.add_style("Codigo", 1)  # 1 = WD_STYLE_TYPE.PARAGRAPH
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.5)
code_style.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
code_style.paragraph_format.space_after = Pt(10)
code_style.paragraph_format.space_before = Pt(6)
code_style.paragraph_format.left_indent = Inches(0.25)
code_style.paragraph_format.line_spacing = 1.0

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")


def write_inline(paragraph, text, bold=False):
    """Escreve texto com **negrito**, `código` e links resolvidos."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            inner = piece[2:-2]
            run = paragraph.add_run(inner.strip("`"))
            run.bold = True
            if inner.startswith("`"):  # **`código`** — negrito e monoespaçado
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
        elif piece.startswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif piece.startswith("["):
            label = piece[1 : piece.index("]")]
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x1A, 0x4B, 0x8C)
        elif piece.startswith("<http"):
            run = paragraph.add_run(piece[1:-1])
            run.font.color.rgb = RGBColor(0x1A, 0x4B, 0x8C)
        else:
            run = paragraph.add_run(piece)
        run.bold = run.bold or bold


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    line = lines[i]

    if line.startswith("```"):
        language = line[3:].strip()
        i += 1
        body = []
        while i < len(lines) and not lines[i].startswith("```"):
            body.append(lines[i])
            i += 1
        i += 1

        if language == "mermaid":
            doc.add_picture(IMG, width=Inches(6.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run("Diagrama de classes do back-end")
            run.italic = True
            run.font.size = Pt(9)
        else:
            doc.add_paragraph("\n".join(body), style="Codigo")
        continue

    if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1]) <= set("|-: "):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        header, body = rows[0], rows[2:]
        table = doc.add_table(rows=0, cols=len(header))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for source, is_header in [(header, True)] + [(r, False) for r in body]:
            cells = table.add_row().cells
            for cell, text in zip(cells, source):
                cell.paragraphs[0].text = ""
                write_inline(cell.paragraphs[0], text, bold=is_header)
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9.5)
        doc.add_paragraph()
        continue

    if line.startswith("#"):
        level = len(line) - len(line.lstrip("#"))
        doc.add_heading(line[level:].strip(), level=min(level, 4))
        i += 1
        continue

    if line.startswith("> "):
        body = []
        while i < len(lines) and lines[i].startswith(">"):
            body.append(lines[i].lstrip(">").strip())
            i += 1
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.35)
        write_inline(paragraph, " ".join(body))
        continue

    match = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", line)
    if match:
        indent, marker, text = match.groups()
        # Junta as linhas de continuação do item (recuadas, sem marcador).
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(-|\d+\.)\s", lines[i]) \
                and lines[i].startswith(" "):
            text += " " + lines[i].strip()
            i += 1
        style = "List Number" if marker.endswith(".") else "List Bullet"
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Inches(0.3 + 0.3 * (len(indent) // 2))
        paragraph.paragraph_format.space_after = Pt(4)
        write_inline(paragraph, text)
        continue

    if line.strip() == "":
        i += 1
        continue

    body = []
    while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "```", "|", "> ")) \
            and not re.match(r"^\s*(-|\d+\.)\s", lines[i]):
        body.append(lines[i].strip())
        i += 1
    write_inline(doc.add_paragraph(), " ".join(body))

doc.save(DST)
print("gerado", os.path.relpath(DST))
